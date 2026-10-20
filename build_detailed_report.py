from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(13)
footer=sec.footer.paragraphs[0]; footer.alignment=2
footer.add_run('Research progress report  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,2)
def page(title): doc.add_page_break(); doc.add_heading(title,1)
def table(headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for cell,txt in zip(t.rows[0].cells,headers): cell.text=txt
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row): cell.text=str(txt)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
            pr.append(borders)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DCE6EF' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF')); pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'85'); x.set(qn('w:type'),'dxa'); mar.append(x)
            pr.append(mar)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.size=Pt(10); r.bold=i==0
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_paragraph('Drug Target Affinity Prediction with Uncertainty', 'Title')
doc.add_paragraph('Research rationale methods evidence and next experiments','Subtitle')
p('Progress through 26 September 2026')
h('Purpose and current finding')
p('We are studying whether a graph and protein language model can predict drug–target affinity more reliably when it learns both an affinity estimate and an uncertainty estimate. This report explains why the work is needed, how the data and models were prepared, what the experiments show, and what evidence is still required for a journal submission.')
p('On our restricted Davis test set, the five-model UQ ensemble reduced MSE from 0.3638 to 0.3240 and MAE from 0.3582 to 0.3067 relative to the matched baseline ensemble. Existing bootstrap intervals favor UQ under the tested resampling schemes. These are encouraging within-protocol results, with statistical and implementation qualifications explained in Sections 7 and 10.')
p('The KIBA baseline ensemble is complete. Its MSE is 0.4075 on 19,540 test interactions. The latest available UQ log reaches seed 314 at epoch 120; a final KIBA UQ result is not yet available. Consequently, replication across both datasets has not been established.')
h('Reading guide')
table(['Section','What the reader will learn'],[
('1–2','Scientific motivation and the data problem'),('3–4','Dataset construction and the implemented architecture'),('5–6','Training protocol and uncertainty evaluation'),('7–8','Davis evidence and KIBA progress'),('9–10','Operational workflow and limits of the current claims'),('11–12','Next experiments and supporting references')])
p('HGTA and UQ-HGTA are project labels here. External references support the scientific methods; the numerical experiment results are project observations, not results attributed to those papers.')

page('1 Scientific motivation and research question')
p('Drug–target affinity prediction assigns a continuous interaction score to a small molecule and a protein. Such predictions can prioritize pairs for experimental measurement. They do not establish therapeutic efficacy, selectivity in cells, safety, or clinical benefit. Our immediate scientific target is accurate regression and reliable predictive uncertainty on kinase benchmarks.')
p('DeepDTA demonstrated a sequence-based approach using learned drug and protein representations [1]. GraphDTA investigated molecular graphs for the drug encoder [2]. These studies provide relevant methodological foundations, but neither constitutes a matched external baseline in our current experiments. Our “baseline” is journal_hgta.py, a project implementation sharing its encoder and fusion design with uq_hgta.py.')
h('Why graphs and protein representations are used')
p('Atoms and bonds supply an explicit representation of molecular connectivity. The graph encoder can use bond features while propagating information among connected atoms. Protein language models such as ESM-2 provide sequence representations learned from large protein collections [5]. Freezing this encoder makes supervised training smaller and allows its outputs to be cached. Whether this choice improves our task must still be tested by ablation; prior success is motivation, not proof.')
h('Why uncertainty is part of the question')
p('A point estimate alone does not indicate whether a prediction is dependable. Input-dependent variance learning and deep ensembles offer established methods for representing predictive uncertainty [6,7]. Our hypothesis is that a probabilistic objective may improve prediction and provide useful intervals. This is testable: lower error does not automatically imply calibrated uncertainty, and wide intervals can attain high coverage without being useful.')
h('Questions the experiments can answer')
p('The first question is whether UQ improves error and ranking compared with the matched MSE-trained model on fixed restricted benchmarks. The second is whether its predicted intervals achieve useful coverage at reasonable width. The third, still open, is whether any benefit persists across seeds, datasets and harder splits containing unseen drugs or targets.')
p('Davis uses binding measurements associated with a kinase inhibitor panel [3]. In the benchmark processing, nanomolar Kd values are converted to pKd, so larger values indicate stronger binding. KIBA is an integrated bioactivity score built from heterogeneous measurements [4]; it must remain on its supplied scale. Davis and KIBA MSE values therefore should not be compared as if they measured the same units.')

page('2 Data integrity and why the experiment changed')
p('The early CSV workflow revealed 970 duplicate canonical input groups in Davis training and 241 in test. Of these, 451 training and 116 test groups had differing labels. There were also 595 canonical input pairs shared between those train and test files. These counts are project audit observations from the earlier input files, not universal properties asserted for every Davis release.')
p('A duplicate model input means the same canonical molecular representation and the same protein sequence. Raw biological identifiers are a different concept: two source IDs may map to identical model inputs. A duplicate within one partition is not by itself proof of train–test leakage. An identical input appearing in training and test creates a direct overlap concern. We therefore distinguish duplicate records, input ambiguity and cross-partition overlap.')
h('Why median aggregation was abandoned')
p('Median aggregation was used in an early exploratory script to make training run, but the conflicting labels were not verified biological replicates. Source reconstruction indicated multiple target IDs sharing stored sequences. A sequence-only model cannot infer an omitted mutation or assay identity from an identical input. Averaging those labels would define a new target quantity without biological justification. The restricted runs instead exclude ambiguous inputs and preserve identifiers.')
h('What the correction establishes')
p('The current launcher rejects duplicate raw ID pairs, repeated canonical input pairs and repeated pair indices across all saved partitions. This establishes exact-input separation for the loaded CSVs. It does not establish independence of similar molecules, homologous proteins, pretraining data, or assay provenance. “Exact-input overlap checked” is more precise than an unrestricted claim that the benchmark is leakage-free.')
h('Why the test set remains limited')
p('Both restricted benchmarks principally evaluate held-out interactions in a pair-split setting. A drug or protein can occur in both training and test when paired with different partners. This is a valid interpolation question, but it does not demonstrate prediction for an unseen drug, scaffold or target family. Separate cold-drug, cold-target and similarity-aware experiments are needed for those claims.')
p('Filtering changes the evaluation population. Removing ambiguous rows can alter target families, affinity distributions and difficulty. We must publish exclusions and describe these experiments as restricted benchmarks. It would be misleading to compare their scores directly with published full-dataset scores and claim superiority.')

page('3 How the two benchmarks were prepared')
h('Davis reconstruction and fixed source folds')
p('The reconstructed benchmark contains 68 drug IDs and 442 target IDs, with 30,056 matrix entries. This processed benchmark is distinct from the original experimental paper’s 72-inhibitor panel [3]. The project audit found only 379 unique stored protein sequences. We retained target identities whose sequence mapped to exactly one source target ID. This excluded 5,508 interactions and retained 361 target IDs.')
p('The local manifest pins GraphDTA source revision 6162245c6776684b391d30610850fe7077478435 and records validation fold 0. The restricted Davis training, validation and test memberships are saved explicitly. Its held-out validation set is derived from source folds, not regenerated for each training seed. Source conventions are documented in the GraphDTA repository [8].')
h('KIBA export and input filtering')
p('The KIBA exporter reads ordered ligand/protein dictionaries, the affinity matrix and fold-setting-1 indices. It exports observed entries with raw drug IDs, target IDs and matrix coordinates. KIBA scores are not log-transformed. The preparer then takes a fixed 15% validation sample from the training pool using seed 20260915. Model training seeds do not change that split.')
p('The initial raw-ID preparation did not eliminate all repeated model inputs. A subsequent filter retained only canonical molecule–sequence pairs occurring exactly once across the prepared train, validation and test files. It removed all rows in each repeated-input group, even if labels matched. This is a conservative engineering choice, not proof that every removed record was erroneous. It differs from the target-wide Davis exclusion rule.')
table(['Dataset','Train','Validation','Test','Excluded'],[
('Restricted Davis','16,475','4,042','4,031','5,508'),('Restricted KIBA','82,979','14,665','19,540','1,070 of 118,254')])
p('KIBA retained 117,184 interactions; the final repeated-input filter removed approximately 0.90% of the 118,254 rows entering that stage. Earlier preparation can also deduplicate or remove overlaps; a complete provenance audit must account for every stage. Source commit and raw-file hashes still need to be recovered for the KIBA export because the supplied clone command followed the current branch.')
h('Why files and hashes matter')
p('Each final directory contains train.csv, validation.csv, test.csv and subset_manifest.json. SHA-256 hashes detect byte changes. IDs support traceability; pair_index is an internal unique index and is not necessarily the original matrix index after adaptation. A hash proves file identity, not biological correctness.')

page('4 The implemented model architecture')
p('The local implementations use the same feature extractor and fusion structure. The baseline predicts one value; UQ predicts two. Architecture details below come from journal_hgta.py and uq_hgta.py. The run-specific stored trainer hashes should be compared with these files before the final manuscript is submitted.')
table(['Stage','Implementation','Reason and tradeoff'],[
('Drug input','12 atom features and 7 bond features','Retains connectivity and bond attributes; no 3D pose'),
('Graph encoder','3 residual GINE layers, width 256','Learns local chemical context; finite neighborhood depth'),
('Protein encoder','Frozen esm2_t12_35M_UR50D','Reuses sequence representations; no task fine-tuning'),
('Sequence processing','Chunks of 1,022 residues; at most 512 pooled tokens','Bounds memory; loses cross-chunk context and residue resolution'),
('Fusion','2 bidirectional attention blocks, 8 heads','Conditions drug features on protein features and vice versa'),
('Readout','Masked mean pooling; four-vector concatenation','Produces fixed-size pair representation'),
('Prediction head','1024 to 512 to 128 to 1 or 2','Outputs affinity mean, optionally log variance')])
p('Atom features include scaled atomic number, degree, charge, hydrogen count, valence, mass, aromaticity, ring status, hybridization and chirality. Bond features encode bond type, conjugation, ring membership and stereo. Linear projections map both to width 256. GINE incorporates edge attributes during message passing; the related graph architecture literature is described in [9]. The code trains this graph encoder from scratch.')
p('Protein embeddings are generated in inference mode, with special tokens removed. Long-sequence chunk embeddings are concatenated and adaptively average-pooled when they exceed 512 tokens. Thus “atom–residue attention” is approximate for long proteins: pooled positions can represent several residues. No docking pose or experimentally observed residue contact is supplied.')
p('After attention, padding-aware means produce drug and protein vectors. Their concatenation, elementwise product and absolute difference form the 1,024-dimensional readout. Attention weights should not be interpreted as validated binding sites. A biological interpretation would require separate contact or perturbation evidence.')

page('5 Training objectives and checkpoint selection')
h('Baseline loss')
p('The baseline minimizes mean squared error: square the difference between each predicted affinity and its label, then average over the batch. Large errors receive greater penalty. This gives a direct point-prediction comparator for the UQ objective.')
h('Uncertainty loss')
p('UQ predicts a mean and a log variance. Its Gaussian negative log-likelihood is one half of the average of the log total variance plus squared error divided by total variance. The constant Gaussian normalization term is omitted during training. The variance term penalizes inflating uncertainty indefinitely, while the inverse-variance weighting changes how residuals contribute to learning [6,7].')
p('The code clamps log variance between −8 and 4, adds a variance floor of 0.0001 and caps total variance at 100. Replicate uncertainty is disabled in these runs. The launcher sets affinity_std to zero and measurement_count to one; these values are API compatibility fields, not experimental estimates of assay noise.')
p('A negative printed UQ training loss is mathematically possible for a continuous Gaussian likelihood, particularly when the normalization constant is omitted. It is not a negative MSE. The history column is currently named train_mse even in UQ; this name is misleading and must be relabeled when plotting or exporting tables.')
table(['Training control','Current setting'],[
('Model seeds','4221, 17, 71, 314, 2025'),('Optimizer','AdamW; learning rate 0.0002; weight decay 0.00001'),('Batch size and dropout','16 and 0.20'),('Budget and early stopping','Maximum 300 epochs; patience 35'),('Learning rate scheduler','Validation MSE; factor 0.5; patience 10; minimum 0.000001'),('Gradient clipping','Norm limit 1.0'),('Checkpoint criterion','Minimum validation MSE for both models')])
p('Each seed trains on the same partition. After training ends, its validation-best checkpoint is loaded and test predictions are produced. Validation selection reduces direct test selection bias [10]. However, repeated human inspection of test results can still influence later development; the study should be described as exploratory development followed by fixed-setting replication, not as formally preregistered.')
p('Using the same seed labels matches the experimental schedule but does not ensure identical initial tensors between one-output and two-output models. Equal maximum epochs also do not mean equal computation, because early stopping occurs at different epochs.')

page('6 How accuracy and uncertainty are evaluated')
table(['Measure','Meaning and interpretation'],[
('MSE and RMSE','Squared error average and its square root; lower is better'),('MAE','Average absolute error; lower is better'),('R squared','Improvement over predicting the test-label mean; can be negative'),('Pearson correlation','Linear association; does not guarantee correct scale or offset'),('Concordance index','Pairwise ordering agreement; label ties excluded, prediction ties receive half credit'),('Coverage and width','Fraction inside an interval and its average size; both are necessary')])
h('Why both individual runs and ensembles are reported')
p('Individual-run mean and standard deviation describe variation across the five fitted models. Ensemble performance is calculated after averaging their predictions for each pair. Averaging MSE values is not the same as measuring the MSE of averaged predictions. A single-seed summary has undefined sample standard deviation, which explains the NaN values printed after restart runs.')
p('For UQ, total ensemble variance is computed as the mean of member variances plus the variance of member means, with division by the number of members. This is the equal-weight mixture variance identity. We call the first component predicted within-model variance and the second between-model disagreement. They are useful proxies for aleatoric and epistemic uncertainty, but neither is a direct experimental measurement of those quantities [6,7].')
h('Intervals and their limitations')
p('The existing ensemble script uses normal intervals with the ensemble mean and total variance. Their approximate 90% and 95% multipliers are 1.645 and 1.960. A mixture of Gaussians need not itself be Gaussian, so these are approximate intervals. Likewise GaussianApproxNLL evaluates a moment-matched Gaussian, not the exact mixture density.')
p('Coverage near the nominal level on one test set does not prove calibration generally. Large intervals can conceal poor precision, and average coverage can mask undercoverage for certain drugs or target groups. Future evaluation should include coverage across nominal levels, interval width, groupwise coverage and exact mixture likelihood.')
p('The prediction routine exports the exponential of half the log variance, whereas training adds the small variance floor. That train–inference mismatch should be reconciled in a versioned implementation before definitive uncertainty claims. A calibration scale must be fitted on a separate calibration partition excluded from training and checkpoint selection; it must not be fitted on test labels.')

page('7 Davis results and the statistical interpretation')
table(['Metric','Baseline ensemble','UQ ensemble'],[
('MSE','0.363824','0.323969'),('RMSE','0.603179','0.569182'),('MAE','0.358197','0.306720'),('R squared','0.434635','0.496569'),('Pearson','0.679793','0.706198'),('Concordance index','0.822893','0.850846')])
p('On 4,031 test pairs, ensemble MSE decreased by 0.039856, or 10.95%, and MAE by 0.051477, or 14.37%. Individual-run MSE was 0.4602 ± 0.0750 for baseline and 0.4044 ± 0.0331 for UQ. UQ lowered per-seed MSE in three of five matched seed comparisons, so the benefit is not uniform across every initialization.')
table(['Resampling scheme','MSE difference 95% interval','MAE difference 95% interval'],[
('Pairs','−0.05478 to −0.02533','−0.05848 to −0.04427'),('Drugs, 68 clusters','−0.06373 to −0.01766','−0.06544 to −0.03750'),('Targets, 361 clusters','−0.05607 to −0.02408','−0.05910 to −0.04367')])
p('These are percentile bootstrap intervals from 10,000 draws using RNG seed 20260915, holding the fitted ensembles fixed. Negative differences favor UQ. Pair resampling treats interactions as independent; one-way cluster resampling retains all rows of a sampled drug or target and accounts for that particular dependency. The intervals provide evidence of improvement under those assumptions.')
p('Two corrections to earlier interpretations are necessary. First, the script’s p_two_sided is twice an empirical bootstrap sign-tail fraction, not a separately validated null-based test. Zero tail counts do not establish a true p-value below 0.0002. The report therefore emphasizes intervals and does not reproduce those values as definitive significance tests. Second, separate drug and target bootstraps do not jointly handle crossed drug–target dependence and are not guaranteed conservative. Multiway resampling is a relevant next analysis [11].')
p('UQ ensemble coverage was 94.10% at nominal 90% and 95.36% at nominal 95%, with mean widths 1.566 and 1.866. The mean single-model 95% coverage was only 92.17%. Ensemble uncertainty improves aggregate coverage here, but this is empirical evidence on this subset rather than a calibration guarantee.')

page('8 KIBA progress and what the logs show')
table(['Completed baseline ensemble measure','Value'],[
('Test interactions','19,540'),('MSE','0.407549'),('RMSE','0.638395'),('MAE','0.413105'),('R squared','0.395422'),('Pearson','0.664982'),('Concordance index','0.768754')])
p('The five-seed KIBA baseline is complete, including reruns after a power interruption. Seed 2025 achieved test MSE 0.5453. Seed 314 reached test MSE 0.2231 after running to the 300-epoch limit. Its displayed validation MSE declined to 0.2347 at epoch 300. The substantial difference between seeds is a central result requiring investigation, not simply a reason to discard weaker runs.')
p('The equal-weight baseline ensemble has higher test error than seed 314 alone. Ensembles can reduce variance but need not outperform their strongest member. Selecting seed 314 because of its observed test score would bias evaluation. The original five-seed ensemble should remain the reported planned comparison; any new ensemble selection rule must be developed using validation data and documented as a new experiment.')
h('Status of the UQ experiment')
p('The latest supplied log shows seeds 4221, 17 and 71 followed by seed 314 through epoch 120, where validation MSE is 0.3370. Seeds 4221 and 71 show declining training objective but deteriorating validation MSE. This suggests poor generalization or optimization instability for those runs; the logs alone cannot establish the cause. Seed 2025 and the final UQ summary have not been provided.')
p('We should finish the current protocol before interpreting the overall comparison. The scientific question is whether UQ adds benefit under matched conditions; a null or negative result on KIBA is informative and must be reported. It may show that the Davis gain is dataset-dependent or sensitive to optimization.')
h('Why KIBA is useful and what it cannot establish')
p('KIBA adds a larger benchmark with a different label construction and a much larger compound set. Repeating the pipeline tests whether an observed effect extends beyond one benchmark. It is a second-benchmark replication, not independent laboratory validation: both datasets concern kinase biology, and biological or source overlap must be audited before claiming independent populations [3,4].')

page('9 Process records and reproducibility')
table(['Stage','Project artifact','Purpose'],[
('Davis source reconstruction','reconstruct_davis.py','Recover raw IDs and source fold identity'),('Davis restriction','prepare_davis_subset.py','Exclude target IDs with shared sequences'),('KIBA export','export_graphdta_kiba_csv.py','Export raw IDs, matrix indices and KIBA scores'),('KIBA partitioning','prepare_kiba_locked_split_v2.py','Prepare fixed validation membership'),('KIBA input restriction','make_kiba_input_unambiguous.py','Remove repeated canonical input groups'),('Training entry point','run_davis_subset.py and run_kiba_subset.py','Verify files and inject saved partitions'),('Model implementations','journal_hgta.py and uq_hgta.py','Train matched baseline and UQ'),('Analysis','ensemble_hgta.py and bootstrap scripts','Combine predictions and compare errors')])
p('The KIBA launcher is an alias for the fixed-partition Davis launcher. The earlier failures were schema mismatches: the preparer wrote val.csv and audit_manifest.json, whereas the launcher required validation.csv, a sha256 mapping and pair_index plus feature columns. Those repairs changed the data interface. The repeated-input error was a substantive audit finding, handled by the restriction step.')
p('For each completed seed, preserve best.pt, history.csv, metrics.json and test_predictions.csv. Preserve the top-level fixed_split_manifest.json, manifest.json and membership CSVs from every original and restarted run. The combined baseline folder contains selected seed directories; it is not by itself a complete provenance archive.')
p('A missing test_predictions.csv indicates that the expected completed artifact is unavailable, but does not alone prove training failed: files can be copied incompletely. Before rerunning, inspect the original output location. The current checkpoint stores model weights, epoch and validation metrics; it does not store optimizer, scheduler, scaler and RNG states needed for exact interrupted-run continuation. Restarting that seed from the beginning is a documented rerun, not checkpoint resumption.')
p('For the next analysis, verify exactly five intended seed directories, unique IDs, finite labels and predictions, and exact equality of the 19,540 pair sets across all seeds and both models. Keep backup directories outside the seed glob. Archive environment versions and model revisions, and record the actual source hashes on the Linux training host.')

page('10 Claims that need qualification before submission')
h('Statistical implementation')
p('The current bootstrap reader uses seed_*/test_predictions.csv, so an extra backup directory matching that pattern could enter the ensemble. Its final UQ–baseline merge is an inner join and can silently evaluate only intersecting pairs. Both issues require explicit seed lists and exact set-equality checks before final statistics. The Davis table reflects the reported 4,031-pair analysis, but the script should be audited and rerun before manuscript submission.')
p('Bootstrap intervals hold the five fitted models fixed and do not quantify variation caused by choosing a new set of training seeds or a new data split. Five seeds provide a limited view of training stability. Drug and target grouping by raw IDs may also differ from grouping by canonical chemistry or protein similarity; the grouping unit must match the scientific claim.')
h('Data and reproducibility')
p('The KIBA pipeline filters after partitioning and can remove test rows. The result is a restricted test population. Source IDs are preserved in version 2 of the preparer, but a full exclusion ledger is still needed. For proteins with related sequences and molecules with similar scaffolds, exact input uniqueness does not remove statistical dependence or establish out-of-domain performance.')
p('The protein cache key contains the model name and sequence, but not all token-pooling settings or a pinned model revision. Reusing a cache after changing preprocessing could therefore give inconsistent features. Pin the pretrained-model revision, include preprocessing in cache identity and verify cache metadata in subsequent experiments.')
h('Novelty and the journal claim')
p('Graph encoding, protein language representations, attention, heteroscedastic regression and deep ensembles are established ideas [2,5–7,9]. Combining them does not alone establish methodological novelty. The present evidence motivates a reproducibility and uncertainty study. A stronger contribution could demonstrate where uncertainty succeeds or fails under identity ambiguity and distribution shift, with ablations and matched external baselines.')
p('The existing results should not be called SOTA: the evaluation populations differ from standard full benchmarks, external contemporary methods have not been retrained on these exact splits, and no comprehensive current literature comparison has been completed. Neither formal preregistration nor universal leakage elimination has been demonstrated.')

page('11 Next experiments and presentation narrative')
h('Immediate completion')
p('Finish all five KIBA UQ runs with the current settings. Preserve the configuration and outputs. After checking prediction completeness and pair alignment, compute both individual-model summaries and the equal-weight ensemble. Compare it with results_kiba_baseline_complete, documenting the restarted seeds and retaining their original manifests.')
h('Statistical and uncertainty validation')
p('Revise the analysis checks described in Section 10, retain effect sizes and confidence intervals, and investigate a method appropriate for crossed drug–target dependence [11]. For uncertainty, report coverage together with width, exact Gaussian-mixture likelihood and group-level failures. Any post-hoc scale or conformal calibration should use a dedicated calibration set and a newly specified evaluation protocol.')
h('Experiments that test the explanation')
p('Ablations should isolate the protein encoder, cross-attention, variance head and ensemble contribution. Compare single baseline, baseline ensemble, single UQ and UQ ensemble under equal training and selection rules. Use validation-only investigation to examine seed instability, learning-rate schedules and stopping budgets. Do not extend only the seeds that look favorable on test.')
p('For generalization, construct cold-drug and cold-target splits, followed by scaffold- or protein-similarity-aware splits if supported by sample size. Refit all preprocessing and selection steps appropriately. Retrain representative published methods on the same partitions and budget before making comparative claims. Record runtime and parameter counts so any gain can be assessed against cost.')
h('Suggested explanation during a presentation')
p('“We started by asking whether adding uncertainty could improve drug–target affinity prediction. During data inspection, we found that raw identities and the inputs visible to the model were not always equivalent. We therefore reconstructed source identities and built explicit restricted splits. Using the same graph and protein encoders, we compared an MSE baseline with a probabilistic head across five seeds. On restricted Davis, the UQ ensemble lowered MSE by about 11%, with bootstrap intervals favoring the improvement. KIBA is our next benchmark: its baseline is complete, but UQ is still running. Our next objective is to verify the gain, explain seed instability and test whether uncertainty remains useful for unfamiliar drugs and targets.”')
p('The defensible present conclusion is an observed Davis improvement under the tested restricted protocol, with encouraging resampling evidence and unresolved replication and calibration questions. This defines a concrete research direction without assuming its final outcome.')

page('12 References and evidence records')
refs=[
'[1] Öztürk H, Özgür A, Ozkirimli E. DeepDTA: deep drug–target binding affinity prediction. Bioinformatics. 2018;34(17):i821–i829. https://doi.org/10.1093/bioinformatics/bty593',
'[2] Nguyen T, Le H, Quinn TP, Nguyen T, Le TD, Venkatesh S. GraphDTA: predicting drug–target binding affinity with graph neural networks. Bioinformatics. 2021;37(8):1140–1147. https://doi.org/10.1093/bioinformatics/btaa921',
'[3] Davis MI, Hunt JP, Herrgard S, et al. Comprehensive analysis of kinase inhibitor selectivity. Nature Biotechnology. 2011;29:1046–1051. https://doi.org/10.1038/nbt.1990',
'[4] Tang J, Szwajda A, Shakyawar S, et al. Making sense of large-scale kinase inhibitor bioactivity data sets: a comparative and integrative analysis. Journal of Chemical Information and Modeling. 2014;54(3):735–743. https://doi.org/10.1021/ci400709d',
'[5] Lin Z, Akin H, Rao R, et al. Evolutionary-scale prediction of atomic-level protein structure with a language model. Science. 2023;379:1123–1130. https://doi.org/10.1126/science.ade2574',
'[6] Kendall A, Gal Y. What uncertainties do we need in Bayesian deep learning for computer vision? Advances in Neural Information Processing Systems 30. 2017. https://arxiv.org/abs/1703.04977',
'[7] Lakshminarayanan B, Pritzel A, Blundell C. Simple and scalable predictive uncertainty estimation using deep ensembles. Advances in Neural Information Processing Systems 30. 2017. https://arxiv.org/abs/1612.01474',
'[8] Nguyen T and contributors. GraphDTA repository, README and create_data.py. https://github.com/thinng/GraphDTA. Davis source revision recorded in this project: 6162245c6776684b391d30610850fe7077478435.',
'[9] Hu W, Liu B, Gomes J, et al. Strategies for pre-training graph neural networks. ICLR. 2020. https://arxiv.org/abs/1905.12265',
'[10] Cawley GC, Talbot NLC. On over-fitting in model selection and subsequent selection bias in performance evaluation. Journal of Machine Learning Research. 2010;11:2079–2107. https://www.jmlr.org/papers/v11/cawley10a.html',
'[11] Owen AB, Eckles D. Bootstrapping data arrays of arbitrary order. Annals of Applied Statistics. 2012;6(3):895–927. https://doi.org/10.1214/12-AOAS547']
for r in refs: p(r)
h('Project evidence')
p('Methods: local journal_hgta.py, uq_hgta.py, run_davis_subset.py and paired_bootstrap_ensemble_v2.py inspected for this report. Davis membership: davis_unambiguous_v1/subset_manifest.json. Numerical results: reported restricted Davis summaries, ensemble/bootstrap outputs and KIBA training outputs through 26 September 2026. Final publication tables should be regenerated from archived prediction files and tied to the training-host hashes.')
out=ROOT/'HGTA_Detailed_Research_Report.docx'; doc.save(out)
print(out)
